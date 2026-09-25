import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_user_guide():
    doc = Document()

    # 1. Page Geometry matching GWC Academic Template
    # Letter 8.5" x 11", Margins: Top=1.5", Bottom=1.0", Left=1.5", Right=1.0"
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.5)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    # 2. Configure Default 'Normal' Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.space_before = Pt(0)

    # Helper Functions
    def add_title_p(text, size=Pt(14), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(6)):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = size
        run.font.bold = bold
        return p

    def add_section_h1(title_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(title_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        return p

    def add_subheading_h2(subtitle_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(subtitle_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.italic = True
        return p

    def add_step_p(step_label, step_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        r_step = p.add_run(step_label + " ")
        r_step.font.name = 'Times New Roman'
        r_step.font.size = Pt(12)
        r_step.font.bold = True
        r_txt = p.add_run(step_text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(12)
        return p

    def add_checklist_item(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.5
        r_box = p.add_run("☑ ")
        r_box.font.name = 'Times New Roman'
        r_box.font.size = Pt(12)
        r_box.font.bold = True
        r_txt = p.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(12)
        return p

    def add_bullet_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.5
        r_bullet = p.add_run("•  ")
        r_bullet.font.name = 'Times New Roman'
        r_bullet.font.size = Pt(12)
        r_bullet.font.bold = True
        r_txt = p.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(12)
        return p

    def add_body_p(text, italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.italic = italic
        return p

    def add_screenshot_figure(image_filename, figure_title, width=Inches(5.5)):
        img_path = os.path.join("screenshots", image_filename)
        p_ph = doc.add_paragraph()
        p_ph.paragraph_format.space_before = Pt(6)
        p_ph.paragraph_format.space_after = Pt(2)
        p_ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ph = p_ph.add_run(f"[Insert Screenshot – {figure_title}]")
        r_ph.font.name = 'Times New Roman'
        r_ph.font.size = Pt(11)
        r_ph.font.bold = True
        r_ph.font.color.rgb = RGBColor(70, 70, 70)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=width)

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run(f"Figure: {figure_title}")
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(10)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(80, 80, 80)
        else:
            p_err = doc.add_paragraph(f"(Image file not found: {img_path})")
            p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_err.runs[0].font.italic = True

    # ==========================================
    # DOCUMENT HEADER & TITLE
    # ==========================================
    add_title_p("BOLINAO SCHOOL OF FISHERIES", size=Pt(14), bold=True)
    add_title_p("STUDENT E-VOTING SYSTEM", size=Pt(13), bold=True)
    add_title_p("APPENDIX H – USER MANUAL", size=Pt(14), bold=True, space_after=Pt(12))

    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_intro.paragraph_format.space_after = Pt(14)
    p_intro.paragraph_format.line_spacing = 1.5
    r_intro = p_intro.add_run(
        "This user manual provides step-by-step instructions for setting up, administering, and using the Bolinao School of Fisheries Student E-Voting System. The system provides secure role-based access control, DepEd Learner Reference Number (LRN) voter authentication, automated DepEd School Form 1 (SF1) roster importation, digital ballot casting with real-time auditability, and instant canvassing."
    )
    r_intro.font.name = 'Times New Roman'
    r_intro.font.size = Pt(12)

    # ==========================================
    # 1. SYSTEM REQUIREMENTS
    # ==========================================
    add_section_h1("1. SYSTEM REQUIREMENTS")
    add_body_p("Before deploying and using the system, prepare the following hardware and software components:")
    add_bullet_p("Server Host or Local Precinct Computer (Node.js v20.x or higher runtime)")
    add_bullet_p("Client Voting Terminals / Kiosks (Desktop PC, Laptop, or Tablet)")
    add_bullet_p("Google Chrome or Microsoft Edge (Latest version recommended)")
    add_bullet_p("Local Area Network (LAN) router or active Internet connection")
    add_bullet_p("Official DepEd School Form 1 (SF1) student masterlist spreadsheet (.xls or .xlsx)")
    add_bullet_p("Registered System Accounts (Administrator, Teacher/Proctor, and Student Voters)")
    add_body_p("Make sure the host server is active and client terminals can communicate over the network before commencing election operations.", italic=True)

    # ==========================================
    # 2. STARTING THE SYSTEM
    # ==========================================
    add_section_h1("2. STARTING THE SYSTEM")
    add_step_p("Step 1:", "Turn on the server computer and start the CivicFlow backend service.")
    add_step_p("Step 2:", "Power on the client voting terminals in the designated polling room.")
    add_step_p("Step 3:", "Connect each workstation to the election Local Area Network or Wi-Fi.")
    add_step_p("Step 4:", "Open Google Chrome or Microsoft Edge browser.")
    add_step_p("Step 5:", "Navigate to the designated application URL (e.g., http://localhost:3000 or the school portal address).")
    add_step_p("Step 6:", "The Bolinao School of Fisheries Student E-Voting System Login Page will appear.")
    add_screenshot_figure("02_login_desktop.png", "System Login Page (Desktop View)", width=Inches(5.5))

    # ==========================================
    # 3. LOGGING IN
    # ==========================================
    add_section_h1("3. LOGGING IN")
    add_body_p("Authorized users access the system using their assigned credentials based on role.")
    add_subheading_h2("A. Student Voter Login Procedure")
    add_step_p("Step 1:", "Enter your 12-digit Learner Reference Number (LRN) into the Student Number field.")
    add_step_p("Step 2:", "Enter your personal account password.")
    add_step_p("Step 3:", "Click the 'Sign In to Ballot' button.")
    add_step_p("Step 4:", "The system will authenticate your session and automatically load the Active Ballot Page.")

    add_subheading_h2("B. First-Time Password Setup for Imported Students")
    add_body_p("Students imported via the DepEd SF1 roster without passwords will be prompted upon their first login to establish credentials:")
    add_step_p("Step 1:", "Enter your Student LRN into the login field and click Sign In.")
    add_step_p("Step 2:", "The 'First-Time User Password Setup' window will appear displaying your full name.")
    add_step_p("Step 3:", "Enter your desired password in the 'New Password' field (minimum 6 characters).")
    add_step_p("Step 4:", "Re-type your password in the 'Confirm Password' field.")
    add_step_p("Step 5:", "Click 'Create Password & Continue' to finalize your credentials and proceed to the ballot.")

    add_subheading_h2("C. Administrator and Faculty Login")
    add_step_p("Step 1:", "Enter your administrative username (e.g., ADMIN).")
    add_step_p("Step 2:", "Enter your secure administrative password.")
    add_step_p("Step 3:", "Click 'Sign In'.")
    add_step_p("Step 4:", "The system will verify administrative authorization and load the Administrator Dashboard.")

    add_subheading_h2("D. How to Cast Your Ballot Guidance")
    add_body_p("Voters unfamiliar with digital voting can click the 'How to Vote' link on the login card to review the three-step voting overview before signing in.")
    add_screenshot_figure("03_modal_how_to_vote_mobile.png", "How to Cast Your Ballot Informational Modal", width=Inches(3.2))

    add_subheading_h2("If Login Fails:")
    add_bullet_p("Verify that the Student LRN contains exactly 12 digits without spaces or dashes.")
    add_bullet_p("Check that the keyboard Caps Lock key is turned off.")
    add_bullet_p("Check with the precinct poll clerk if your name is listed on the official DepEd SF1 enrollment roster.")
    add_bullet_p("If your password is forgotten, request a password reset from the Election Administrator.")

    # ==========================================
    # 4. REGISTERING A USER & BULK IMPORT
    # ==========================================
    add_section_h1("4. REGISTERING A USER AND BULK SF1 IMPORT")
    add_body_p("Only authorized administrators can register voter accounts and manage system credentials.")
    add_subheading_h2("A. Manual User Creation")
    add_step_p("Step 1:", "Log in as an Administrator.")
    add_step_p("Step 2:", "Click 'Users' on the navigation drawer menu.")
    add_step_p("Step 3:", "Click the '+ Add User' button.")
    add_step_p("Step 4:", "Enter the user's details: Student Number (LRN), Full Name, Grade Level (7–12), Section (e.g., Glassfish, Sailfish), and designated Room.")
    add_step_p("Step 5:", "Select the user role: Student, Teacher/Proctor, or Admin.")
    add_step_p("Step 6:", "Verify the information and click 'Save User'.")
    add_step_p("Step 7:", "The user will appear in the active user roster.")

    add_subheading_h2("B. Bulk Importing from DepEd School Form 1 (SF1) Excel Files")
    add_body_p("The system provides a built-in parser for rapid importation of complete class rosters directly from standard Department of Education SF1 spreadsheets:")
    add_step_p("Step 1:", "Open the Users tab and click 'Bulk Import'.")
    add_step_p("Step 2:", "Select the 'Upload File' tab or click the drag-and-drop zone.")
    add_step_p("Step 3:", "Choose the target DepEd SF1 spreadsheet (e.g., SF1_2026_Grade 7 (Year I) - GLASSFISH (1).xls).")
    add_step_p("Step 4:", "The system automatically parses student LRNs, full names, year levels, and sections.")
    add_step_p("Step 5:", "Review the validation preview table to ensure no duplicate LRNs exist.")
    add_step_p("Step 6:", "Click 'Import Students'. The system registers all accounts without passwords, allowing students to set their own password on their first login.")
    add_screenshot_figure("admin_bulk_import_modal_mobile.png", "Bulk User Registry Import (DepEd SF1 Parser)", width=Inches(3.4))
    add_screenshot_figure("tab_users_mobile.png", "User Directory and Voter Roster Table", width=Inches(3.4))

    # ==========================================
    # 5. CONFIGURING AN ELECTION EVENT
    # ==========================================
    add_section_h1("5. CONFIGURING AN ELECTION EVENT")
    add_body_p("Before students can vote, an election event must be configured and scheduled.")
    add_step_p("Step 1:", "Navigate to the 'Elections' tab in the Administrator portal.")
    add_step_p("Step 2:", "Click '+ Create Election'.")
    add_step_p("Step 3:", "Enter the Election Title (e.g., Supreme Secondary Learner Government General Election 2026).")
    add_step_p("Step 4:", "Enter a detailed description of the election rules and purpose.")
    add_step_p("Step 5:", "Set the Start Date and Time (when voting officially opens).")
    add_step_p("Step 6:", "Set the End Date and Time (when voting automatically closes).")
    add_step_p("Step 7:", "Select the Election Scope: 'All Students' for general elections, or filter by specific Grade Level, Section, or Voting Room.")
    add_step_p("Step 8:", "Toggle 'Support Party-Lists' if political parties are contesting the election.")
    add_step_p("Step 9:", "Click 'Save Election'. The election status will reflect as 'Upcoming' until the start time is reached.")
    add_screenshot_figure("tab_elections_mobile.png", "Elections Management and Scheduling Tab", width=Inches(3.4))

    # ==========================================
    # 6. MANAGING ELECTORAL POSITIONS AND QUOTAS
    # ==========================================
    add_section_h1("6. MANAGING ELECTORAL POSITIONS AND QUOTAS")
    add_body_p("Positions define the offices contested on the ballot.")
    add_step_p("Step 1:", "Click 'Positions' on the navigation drawer menu.")
    add_step_p("Step 2:", "Select the target election from the dropdown menu.")
    add_step_p("Step 3:", "Click '+ Add Position'.")
    add_step_p("Step 4:", "Enter the official Position Name (e.g., President, Vice President, Secretary, Treasurer, Auditor, Public Information Officer, Grade 7 Representative).")
    add_step_p("Step 5:", "Set the Ballot Order number to arrange positions from highest to lowest executive rank.")
    add_step_p("Step 6:", "Specify the maximum allowable votes (1 for single-winner executive positions; 2 or more for multi-seat councilors).")
    add_step_p("Step 7:", "Click 'Save Position'.")
    add_screenshot_figure("tab_positions_mobile.png", "Positions Management and Ballot Quotas Tab", width=Inches(3.4))

    # ==========================================
    # 7. REGISTERING CANDIDATES AND PARTY-LISTS
    # ==========================================
    add_section_h1("7. REGISTERING CANDIDATES AND PARTY-LISTS")
    add_body_p("Candidate profiles display photos, party affiliations, and platforms on the voting ballot.")
    add_step_p("Step 1:", "Click 'Candidates' on the navigation menu.")
    add_step_p("Step 2:", "Click '+ Add Candidate'.")
    add_step_p("Step 3:", "Select the target Election and Position.")
    add_step_p("Step 4:", "Enter the candidate's Full Name or link their enrolled student account.")
    add_step_p("Step 5:", "Assign their Party-List affiliation (or select 'Independent').")
    add_step_p("Step 6:", "Upload the candidate's portrait photo. Use the built-in Image Cropping tool to zoom, pan, and ensure uniform headshot proportions.")
    add_step_p("Step 7:", "Enter the candidate's platform or advocacy statement.")
    add_step_p("Step 8:", "Click 'Save Candidate'. The candidate profile is published to the ballot.")
    add_screenshot_figure("tab_candidates_mobile.png", "Candidates Management Tab", width=Inches(3.4))
    add_screenshot_figure("student_candidate_modal_mobile.png", "Candidate Profile and Manifesto Modal", width=Inches(3.4))

    # ==========================================
    # 8. CHECKING THE USER AND VOTER ELIGIBILITY
    # ==========================================
    add_section_h1("8. CHECKING THE USER AND VOTER ELIGIBILITY")
    add_body_p("Before students are permitted to cast their ballots, verify that voter registration is complete.")
    add_body_p("Make sure the voter has:")
    add_checklist_item("Active Student Account with verified 12-digit DepEd LRN")
    add_checklist_item("Correct Grade Level, Section, and Room assignment")
    add_checklist_item("Established account password (or ready for first-time setup)")
    add_checklist_item("Unvoted status (Has Voted = False)")
    add_body_p("Voter Readiness Flow:", italic=True)
    add_body_p("Identity Verification → Roster Check → Terminal Login → Active Ballot Access")
    add_body_p("If all checks pass, the student is authorized to cast their vote.")

    # ==========================================
    # 9. CASTING A VOTE (STUDENT VOTER MANUAL)
    # ==========================================
    add_section_h1("9. CASTING A VOTE (STUDENT VOTER MANUAL)")
    add_body_p("This is the standard digital voting procedure for students:")
    add_step_p("Step 1:", "Log in using your Student Number (LRN) and password at the polling terminal.")
    add_step_p("Step 2:", "The Active Election Ballot opens, featuring a live countdown timer indicating the time remaining.")
    add_step_p("Step 3:", "Scroll through the electoral positions in constitutional order.")
    add_step_p("Step 4:", "Click on any candidate's profile card to review their campaign photo, party-list, and full manifesto.")
    add_step_p("Step 5:", "Click the 'Cast Vote' button under your chosen candidate. The selected card will be highlighted with a checkmark badge.")
    add_step_p("Step 6:", "Continue selecting candidates for each position. The system will prevent exceeding the maximum allowable vote quota per position.")
    add_step_p("Step 7:", "Once all selections are made, click 'Review Ballot' at the bottom of the screen.")
    add_step_p("Step 8:", "The Vote Confirmation Modal will appear, displaying a complete summary of all your selections.")
    add_step_p("Step 9:", "Review your ballot carefully. If you wish to change a selection, click 'Cancel' to return to the ballot.")
    add_step_p("Step 10:", "If satisfied, click 'Confirm Vote'.")
    add_step_p("Step 11:", "The system executes an animated ballot drop celebration confirming that your vote has been securely recorded.")
    add_step_p("Step 12:", "A tamper-evident digital receipt confirmation is generated. Click 'Finish & Log Out' to clear the terminal for the next student.")
    add_screenshot_figure("student_vote_flow_mobile.png", "Student Digital Voting Ballot", width=Inches(3.4))
    add_screenshot_figure("student_vote_confirm_modal_mobile.png", "Ballot Review and Confirmation Modal", width=Inches(3.4))

    # ==========================================
    # 10. REAL-TIME MONITORING AND UNVOTED VOTER TRACKING
    # ==========================================
    add_section_h1("10. REAL-TIME MONITORING AND UNVOTED VOTER TRACKING")
    add_body_p("Election officials can monitor turnout and track voter participation in real time:")
    add_step_p("Step 1:", "Log in to the Administrator Dashboard.")
    add_step_p("Step 2:", "Review the live summary tiles displaying Total Users, Total Votes Cast, Total Positions, and Total Candidates.")
    add_step_p("Step 3:", "Inspect turnout percentage progress bars across Grade Levels and Sections.")
    add_step_p("Step 4:", "Navigate to the 'Results' tab and click the 'Unvoted Students' sub-tab.")
    add_step_p("Step 5:", "Use the search bar to locate specific sections or students who have not yet voted.")
    add_step_p("Step 6:", "Coordinate with classroom advisers and proctors to dispatch pending voter batches to the polling booths.")
    add_screenshot_figure("04_admin_dashboard_mobile.png", "Administrator Dashboard Overview", width=Inches(3.4))

    # ==========================================
    # 11. CANVASSING AND VIEWING ELECTION RESULTS
    # ==========================================
    add_section_h1("11. CANVASSING AND VIEWING ELECTION RESULTS")
    add_body_p("The system performs instant, automated canvassing and vote tallying:")
    add_step_p("Step 1:", "Click 'Results' on the main navigation menu.")
    add_step_p("Step 2:", "Select the target election from the election filter dropdown.")
    add_step_p("Step 3:", "Observe the live vote tally per position.")
    add_step_p("Step 4:", "Review the Winner Podium highlighting the leading candidates with Gold, Silver, and Bronze badges.")
    add_step_p("Step 5:", "Examine the interactive Donut Charts illustrating percentage distributions of votes cast.")
    add_step_p("Step 6:", "When the scheduled voting window concludes, the election status transitions to 'Ended', permanently locking the vote counts.")
    add_screenshot_figure("tab_results_mobile.png", "Live Canvassing, Podium Leaderboard, and Results Charts", width=Inches(3.4))

    # ==========================================
    # 12. GENERATING A REPORT
    # ==========================================
    add_section_h1("12. GENERATING A REPORT")
    add_body_p("Official election documents can be generated for institutional records and certification:")
    add_step_p("Step 1:", "Open the Results tab.")
    add_step_p("Step 2:", "Click 'Generate Official Certificate of Canvass'.")
    add_step_p("Step 3:", "Verify the generated report containing total registered voters, total ballots cast, turnout percentage, and candidate rankings.")
    add_step_p("Step 4:", "Click 'Export Report' to save or print the document in PDF or spreadsheet format.")
    add_step_p("Step 5:", "Present the Certificate of Canvass to the School Election Committee, Head Teacher, and Principal for official signing.")

    # ==========================================
    # 13. WHEN A STUDENT CANNOT LOG IN
    # ==========================================
    add_section_h1("13. WHEN A STUDENT CANNOT LOG IN")
    add_body_p("If a student encounters authentication issues:")
    add_step_p("Step 1:", "Check that the student is typing their correct 12-digit DepEd LRN.")
    add_step_p("Step 2:", "Ensure no accidental leading or trailing spaces are present in the student number.")
    add_step_p("Step 3:", "Verify that the keyboard Caps Lock is disabled.")
    add_step_p("Step 4:", "If the student forgot their created password, have an authorized Administrator locate the user in the User Management tab and click 'Reset Password'.")
    add_step_p("Step 5:", "If the system reports 'Account Not Found', verify whether the student was included in the uploaded DepEd SF1 masterlist.")
    add_step_p("Step 6:", "If the student was recently enrolled, the administrator may manually add their record.")

    # ==========================================
    # 14. WHEN A VOTE CANNOT BE CAST
    # ==========================================
    add_section_h1("14. WHEN A VOTE CANNOT BE CAST")
    add_body_p("If the voting portal blocks a student from submitting a ballot:")
    add_step_p("Step 1:", "Check the Election Status. Voting is only permissible during the 'Live' election phase.")
    add_step_p("Step 2:", "Verify the Election Scope. If an election is scoped to a specific Grade Level or Section, ensure the student's profile matches the targeted cohort.")
    add_step_p("Step 3:", "Check the 'Has Voted' status. The system strictly enforces a one-student, one-vote policy. Once a ballot is submitted, subsequent attempts are automatically rejected.")
    add_step_p("Step 4:", "Ensure that selections do not violate ballot quotas (e.g., attempting to select multiple candidates for a single-winner position).")

    # ==========================================
    # 15. MANUAL VERIFICATION AND HELP DESK PROCEDURES
    # ==========================================
    add_section_h1("15. MANUAL VERIFICATION PROCEDURES")
    add_body_p("When automated verification requires human intervention:")
    add_step_p("Step 1:", "The student reports to the Election Committee Polling Precinct Desk.")
    add_step_p("Step 2:", "The poll clerk verifies the student's identity using their physical DepEd School ID card.")
    add_step_p("Step 3:", "The poll clerk cross-references the student's name against the printed DepEd SF1 enrollment roster.")
    add_step_p("Step 4:", "The poll clerk checks the system directory to confirm the voter's active eligibility.")
    add_step_p("Step 5:", "Follow the approved manual verification procedure and assist the voter to an available terminal.")
    add_step_p("Step 6:", "Never create duplicate voter records for a student who has already voted.")

    # ==========================================
    # 16. COMMON PROBLEMS AND SOLUTIONS (TABLE)
    # ==========================================
    add_section_h1("16. COMMON PROBLEMS")
    add_body_p("The following matrix summarizes common operational issues and recommended resolutions:")

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="B0B7C3", sz="4", val="single")

    hdr_cells = table.rows[0].cells
    hdr_cells[0].width = Inches(2.5)
    hdr_cells[1].width = Inches(3.5)
    set_cell_background(hdr_cells[0], "E5E7EB")
    set_cell_background(hdr_cells[1], "E5E7EB")
    set_cell_margins(hdr_cells[0], top=100, bottom=100, left=150, right=150)
    set_cell_margins(hdr_cells[1], top=100, bottom=100, left=150, right=150)

    p0 = hdr_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = p0.add_run("Problem")
    r0.font.name = 'Times New Roman'
    r0.font.bold = True
    r0.font.size = Pt(11)

    p1 = hdr_cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p1.add_run("What to Do")
    r1.font.name = 'Times New Roman'
    r1.font.bold = True
    r1.font.size = Pt(11)

    troubleshooting_data = [
        ("Cannot log in", "Check 12-digit Student LRN and password; turn off Caps Lock."),
        ("First-time login prompt appears", "Enter and confirm a secure new personal password (min. 6 characters)."),
        ("Student account not found", "Check LRN spelling; cross-reference DepEd SF1 enrollment roster; contact administrator."),
        ("System reports 'Already Voted'", "Voters can only cast a ballot once. Duplicate voting is mathematically prohibited."),
        ("Election status shows 'Upcoming'", "Voting has not commenced yet. Await the scheduled election start time."),
        ("Election status shows 'Ended'", "The election deadline has elapsed. The voting window is permanently closed."),
        ("Positions missing on ballot", "Verify election scope settings; some positions are restricted to specific grade levels."),
        ("Candidate photo fails to upload", "Ensure image format is PNG, JPG, or WEBP and file size is below 5 MB. Re-crop using image tool."),
        ("DepEd SF1 bulk import error", "Ensure file is in .xls or .xlsx format with LRN and Student Name column headers."),
        ("Ballot submission is slow or hangs", "Check precinct Wi-Fi or LAN connection; do not close browser until confirmation receipt appears."),
        ("Turnout percentage not updating", "Click the browser refresh button or check connection to the central server."),
        ("Tablet screen distorted or clipped", "Update browser to the latest version of Chrome or Edge; rotate tablet to landscape orientation.")
    ]

    for prob, sol in troubleshooting_data:
        row = table.add_row()
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.5)
        c1.width = Inches(3.5)
        set_cell_margins(c0, top=80, bottom=80, left=140, right=140)
        set_cell_margins(c1, top=80, bottom=80, left=140, right=140)

        p_prob = c0.paragraphs[0]
        p_prob.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_prob.paragraph_format.line_spacing = 1.15
        p_prob.paragraph_format.space_after = Pt(2)
        r_p = p_prob.add_run(prob)
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(10.5)
        r_p.font.bold = True

        p_sol = c1.paragraphs[0]
        p_sol.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_sol.paragraph_format.line_spacing = 1.15
        p_sol.paragraph_format.space_after = Pt(2)
        r_s = p_sol.add_run(sol)
        r_s.font.name = 'Times New Roman'
        r_s.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # 17. SECURITY AND PRIVACY
    # ==========================================
    add_section_h1("17. SECURITY AND PRIVACY")
    add_body_p("All operators, proctors, and voters must strictly observe the following security and data privacy reminders:")
    add_checklist_item("Keep your username, LRN, and account password private at all times.")
    add_checklist_item("Do not use another student's account or LRN to access the system.")
    add_checklist_item("Observe absolute ballot secrecy: selections are cryptographically decoupled from voter identities.")
    add_checklist_item("Adhere to Republic Act No. 10173 (Data Privacy Act of 2012): student masterlists and LRNs must never be exposed or shared outside authorized election activities.")
    add_checklist_item("Do not disclose or copy confidential election audit trails without formal COMELEC approval.")
    add_checklist_item("Only authorized administrative personnel may access the User Management and Canvassing tabs.")
    add_checklist_item("Always log out immediately after casting a vote or completing an administrative task.")

    # ==========================================
    # 18. LOGGING OUT
    # ==========================================
    add_section_h1("18. LOGGING OUT")
    add_body_p("After completing your voting session or administrative duties:")
    add_step_p("Step 1:", "Click your User Profile / Avatar icon in the upper right header.")
    add_step_p("Step 2:", "Click 'Log Out'.")
    add_step_p("Step 3:", "Wait until the system terminates the session token and redirects to the clean Login Page.")
    add_step_p("Step 4:", "Ensure the terminal is reset to the login standby state for the next voter.")
    add_screenshot_figure("02_login_desktop.png", "System Ready for Next Voter (Standby Login Page)", width=Inches(5.5))

    # Save Document
    output_filename = "Bolinao_School_Election_System_USER_GUIDE.docx"
    doc.save(output_filename)
    print(f"Successfully generated User Guide: {output_filename}")

if __name__ == "__main__":
    build_user_guide()
