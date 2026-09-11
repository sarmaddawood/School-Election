import os
import base64
import requests
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()
    doc.add_heading('School Election System - System Design Document', 0)

    # Database Schema
    doc.add_heading('1. Database Schema', level=1)
    
    schema_info = """The system uses an Appwrite database with the following collections:

1. Users: Stores user credentials and profile info.
- Attributes: username, password, fullName, role, studentNumber, yearLevel, section, room, hasSetPassword, photoUrl.

2. Elections: Stores election configuration.
- Attributes: title, description, startsAt, endsAt, scope, scopeValue, hasPartyList, targetGradeLevel, targetSection, targetRoom.

3. Positions: Stores positions available in an election.
- Attributes: electionId, title, name, normalizedName.

4. Candidates: Stores candidate information.
- Attributes: fullName, positionId, electionId, userId, party, partyListId, partyListName, manifesto, photoUrl, voteCount.

5. Votes: Stores individual votes.
- Attributes: userId, voterId, electionId, positionId, candidateId, timestamp, isOfflineImport.

6. Party Lists: Stores party lists for elections.
- Attributes: electionId, name, normalizedName, acronym, logoUrl, advocacy.

7. Branding: Stores school branding settings.
8. Audit Logs: Tracks administrative actions.
9. Offline Ballots: Stores records of imported offline ballots.
"""
    doc.add_paragraph(schema_info)

    def add_mermaid_diagram(title, code, filename):
        doc.add_heading(title, level=1)
        doc.add_paragraph('Mermaid Diagram Code:')
        doc.add_paragraph(code, style='Intense Quote')
        
        try:
            encoded = base64.b64encode(code.encode('utf-8')).decode('utf-8')
            url = f"https://mermaid.ink/img/{encoded}"
            response = requests.get(url)
            if response.status_code == 200:
                with open(filename, 'wb') as f:
                    f.write(response.content)
                doc.add_picture(filename, width=Inches(6))
                os.remove(filename)
        except Exception as e:
            doc.add_paragraph(f"[Error generating image: {e}]")
    
    usecase = """flowchart LR
    Admin([Admin])
    Student([Student])
    
    subgraph School Election System
        M_E[Manage Elections]
        M_C[Manage Candidates]
        M_V[Monitor Votes]
        O_B[Import Offline Ballots]
        L_P[Login to Portal]
        C_V[Cast Vote]
        V_R[View Results]
    end
    
    Admin --- M_E
    Admin --- M_C
    Admin --- M_V
    Admin --- O_B
    
    Student --- L_P
    Student --- C_V
    Student --- V_R"""
    add_mermaid_diagram('2. Use Case Diagram', usecase, 'usecase.png')

    context = """flowchart TD
    A[Admin User] -->|Manages Setup & Data| SES((School Election System))
    S[Student/Voter] -->|Casts Vote| SES
    SES -->|Displays Results| S
    SES -->|Provides Audit Logs| A
    APPWRITE[(Appwrite Database & Storage)] --- SES"""
    add_mermaid_diagram('3. Context Diagram (Level 0 DFD)', context, 'context.png')

    dfd1 = """flowchart TD
    A[Admin] -->|Provides Election Data| P1(1. Manage Elections)
    P1 -->|Stores| DB[(Database)]
    A -->|Provides Candidate Data| P2(2. Manage Candidates)
    P2 -->|Stores| DB
    S[Student] -->|Credentials| P3(3. Authenticate User)
    P3 -->|Reads| DB
    P3 -->|Returns Auth Token| S
    S -->|Selects Candidates| P4(4. Cast Vote)
    P4 -->|Validates & Stores Vote| DB
    P5(5. Tally Results) -->|Reads| DB
    P5 -->|Shows Results| A"""
    add_mermaid_diagram('4. Data Flow Diagram (Level 1)', dfd1, 'dfd1.png')

    erd = """erDiagram
    USERS ||--o{ VOTES : casts
    USERS {
        string id PK
        string studentNumber
        string password
        string fullName
        string role
    }
    ELECTIONS ||--o{ POSITIONS : has
    ELECTIONS ||--o{ PARTY_LISTS : has
    ELECTIONS ||--o{ VOTES : receives
    ELECTIONS {
        string id PK
        string title
        string startsAt
        string endsAt
    }
    POSITIONS ||--o{ CANDIDATES : for
    POSITIONS {
        string id PK
        string electionId FK
        string name
    }
    CANDIDATES ||--o{ VOTES : receives
    CANDIDATES {
        string id PK
        string fullName
        string electionId FK
        string positionId FK
    }
    VOTES {
        string id PK
        string voterId FK
        string electionId FK
        string candidateId FK
    }
    PARTY_LISTS {
        string id PK
        string name
        string electionId FK
    }"""
    add_mermaid_diagram('5. Entity Relationship Diagram (ERD)', erd, 'erd.png')

    flowchart = """flowchart TD
    Start([Start]) --> Login[Student Logs In]
    Login --> Auth{Valid Credentials?}
    Auth -- No --> Err[Show Error] --> Login
    Auth -- Yes --> Dash[Dashboard: Select Election]
    Dash --> Check{Already Voted?}
    Check -- Yes --> End([End - Show Receipt])
    Check -- No --> Ballot[Display Ballot]
    Ballot --> Select[Student Selects Candidates]
    Select --> Confirm{Confirm Choices?}
    Confirm -- No --> Select
    Confirm -- Yes --> Submit[Submit Vote]
    Submit --> DB[(Save to DB)]
    DB --> Receipt[Generate Receipt] --> End"""
    add_mermaid_diagram('6. Flowchart (Voting Process)', flowchart, 'flow.png')

    doc.save('System_Design.docx')
    print("System_Design.docx has been created.")

if __name__ == '__main__':
    create_document()
